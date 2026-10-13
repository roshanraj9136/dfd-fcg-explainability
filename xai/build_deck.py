r"""Build the presentation deck (16:9) from the computed results.

Layout is space-aware: pic_fit() scales each figure to the box it is given and
returns the real bottom edge, so following content can never be pushed off the
slide.  Verified by xai/check_deck.py.
"""
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path("xai/out"); FIG = Path("xai/figs")
R = json.load(open(OUT / "report.json"))
D = json.load(open(OUT / "derived.json"))
RS = json.load(open(OUT / "robustness_summary.json"))
V_R, V_F = R["videos"]["000"], R["videos"]["000_003"]

INK = RGBColor(0x1C, 0x1C, 0x24)
MUTE = RGBColor(0x5A, 0x5F, 0x6E)
BLUE = RGBColor(0x2E, 0x6F, 0xD9)
ORNG = RGBColor(0xD9, 0x62, 0x2E)
GREEN = RGBColor(0x1E, 0x7A, 0x3C)
RED = RGBColor(0xB3, 0x2A, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BAND = RGBColor(0xF2, 0xF4, 0xF8)

SW_IN, SH_IN = 13.333, 7.5
MARGIN = 0.62
CW = SW_IN - 2 * MARGIN          # 12.09 usable width
BOTTOM = 7.18                    # never draw below this

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW_IN), Inches(SH_IN)
BLANK = prs.slide_layouts[6]
pct = lambda v: f"{v:.1f}%"


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    return s


def tb(s, x, y, w, h, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT,
       italic=False, space_after=5, line=1.14):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n") if isinstance(text, str) else text
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line
        r = p.add_run()
        r.text = ("\u2022  " + ln[2:]) if ln.startswith("- ") else ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = "Segoe UI"
    return box


def title(s, main, sub=None, kicker=None):
    bar = s.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(0.085))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    bar.shadow.inherit = False
    y = 0.30
    if kicker:
        tb(s, MARGIN, y, CW, 0.26, kicker.upper(), 11, True, BLUE); y += 0.30
    tb(s, MARGIN, y, CW, 0.52, main, 25, True, INK); y += 0.56
    if sub:
        tb(s, MARGIN, y, CW, 0.34, sub, 13, False, MUTE); y += 0.44
    return y + 0.10


def pic_fit(s, name, x, y, max_w, max_h):
    """Scale to fit inside (max_w, max_h); returns (bottom, width, height)."""
    iw, ih = Image.open(FIG / name).size
    ar = iw / ih
    w = max_w; h = w / ar
    if h > max_h:
        h = max_h; w = h * ar
    xx = x + (max_w - w) / 2          # centre inside the box
    s.shapes.add_picture(str(FIG / name), Inches(xx), Inches(y),
                         width=Inches(w), height=Inches(h))
    return y + h, w, h


def card(s, x, y, w, h, head, value, note=None, accent=BLUE, vs=19):
    bx = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    bx.fill.solid(); bx.fill.fore_color.rgb = BAND
    bx.line.color.rgb = accent; bx.line.width = Pt(1.2)
    bx.shadow.inherit = False
    tb(s, x + .20, y + .14, w - .40, .26, head.upper(), 10, True, MUTE)
    tb(s, x + .20, y + .40, w - .40, .42, value, vs, True, accent)
    if note:
        tb(s, x + .20, y + .88, w - .40, h - 1.0, note, 10.5, False, INK, line=1.18)
    return y + h


def table(s, x, y, w, rows, col_w=None, fs=11, head_fs=11, row_h=0.30):
    nr, nc = len(rows), len(rows[0])
    shp = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w),
                             Inches(row_h * nr))
    t = shp.table
    if col_w:
        tot = sum(col_w)
        for i, c in enumerate(col_w):
            t.columns[i].width = Emu(int(Inches(w) * c / tot))
    for ri, row in enumerate(rows):
        t.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.text = str(val)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.01)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(head_fs if ri == 0 else fs)
                r.font.bold = ri == 0
                r.font.name = "Segoe UI"
                r.font.color.rgb = WHITE if ri == 0 else INK
            c.fill.solid()
            c.fill.fore_color.rgb = (BLUE if ri == 0
                                     else (BAND if ri % 2 else WHITE))
    return y + row_h * nr


# ===================================================================== 1 title
s = slide()
hero = s.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(2.42))
hero.fill.solid(); hero.fill.fore_color.rgb = RGBColor(0x14, 0x2A, 0x54)
hero.line.fill.background(); hero.shadow.inherit = False
tb(s, 0.8, 0.46, 11.8, 0.3, "CVPR 2025  \u00b7  DFD-FCG  \u00b7  EXPLAINABILITY STUDY",
   12, True, RGBColor(0x8F, 0xB6, 0xF5))
tb(s, 0.8, 0.82, 11.8, 0.8, "Why does the deepfake detector say YES or NO?",
   31, True, WHITE)
tb(s, 0.8, 1.70, 11.8, 0.4,
   "Exact spatial vs temporal attribution for a frozen-CLIP video deepfake detector",
   14, False, RGBColor(0xC8, 0xD8, 0xF5))

card(s, MARGIN, 2.72, 3.86, 1.78, "Real sample  000.mp4",
     f"{V_R['verdict']}  \u00b7  P={V_R['prob']['final_mean']:.3f}",
     f"correct \u2713   {V_R['clip_level_accuracy']:.0f}% of clips agree\n"
     f"driver: TEMPORAL {pct(V_R['branch_shapley']['pct_temporal_mean'])}", BLUE, 17)
card(s, 4.72, 2.72, 3.86, 1.78, "Fake sample  000_003.mp4",
     f"{V_F['verdict']}  \u00b7  P={V_F['prob']['final_mean']:.3f}",
     f"correct \u2713   {V_F['clip_level_accuracy']:.0f}% of clips agree\n"
     f"driver: SPATIAL {pct(V_F['branch_shapley']['pct_spatial_mean'])} \u00b7 eyes",
     ORNG, 17)
card(s, 8.85, 2.72, 3.86, 1.78, "Attribution method", "Exact, not estimated",
     f"reproduces every head logit\nto {R['instrumentation_max_logit_error']:.0e}",
     GREEN, 17)

tb(s, MARGIN, 4.80, CW, 1.7,
   ["- Both FaceForensics++ samples shipped with the repository are classified correctly.",
    "- Each decision is split into a spatial vs temporal share, then into facial components, "
    "patches and encoder layers \u2014 all exactly.",
    f"- Stress-tested on the authors' own 7 perturbations \u00d7 3 severities: "
    f"{RS['false_positives']}/21 false positives, {RS['false_negatives']}/21 false negatives."],
   13, False, INK, line=1.34)
tb(s, MARGIN, 6.66, CW, 0.3,
   "Released FFG ViT-L/14 checkpoint  \u00b7  303.4 M parameters, only 0.25 M trained (0.08%)",
   10.5, False, MUTE)

# ===================================================================== 2 background
s = slide()
y = title(s, "What the authors built", "A small side network reads a frozen CLIP encoder",
          "background")
tb(s, MARGIN, y, 6.0, 2.4,
   ["- CLIP ViT-L/14 is FROZEN \u2014 none of the 303 M backbone parameters are trained.",
    "- A side network (0.25 M params) taps the q / k / v / embedding tensors of all 24 layers.",
    "- SPATIAL branch (S): four learned queries \u2014 lips, skin, eyes, nose \u2014 attend over "
    "image patches. This is Facial Component Guidance.",
    "- TEMPORAL branch (T): frame-to-frame affinity per patch, then convolutions; the output "
    "is a 16\u00d716 map.",
    "- Three heads vote: s_head, t_head and fusion a_head."],
   12.5, False, INK, line=1.26)
card(s, 6.85, y, 2.8, 1.32, "trainable", "0.08%", "0.25 M of 303.4 M", GREEN, 18)
card(s, 9.90, y, 2.81, 1.32, "input", "10 frames", "3 s clip, aligned face", BLUE, 18)
tb(s, 6.85, y + 1.50, 5.86, 1.0,
   ["Why this architecture is unusually explainable:",
    "- the score is an explicit mean of three heads;",
    "- y_s and y_t are sums over 24 layers, 4 components and 256 patches."],
   11.5, False, INK, line=1.22)
pic_fit(s, "f1_architecture.png", MARGIN, y + 2.56, CW, BOTTOM - (y + 2.56))

# ===================================================================== 3 method
s = slide()
y = title(s, "How we measure each contribution",
          "No gradients, no surrogate model, no approximation", "method")
tb(s, MARGIN, y, 6.15, 3.5,
   ["The decision is a 2-class softmax, so the entire verdict is one number:",
    "       d  =  logit(FAKE) \u2212 logit(REAL),        P(FAKE) = \u03c3(d)",
    "Each head is a LayerNorm followed by a Linear layer, and its input is a SUM of parts "
    "(24 layers; 4 components; 256 patches).",
    "Because the LayerNorm mean is itself additive over those parts, d splits exactly:",
    "       d  =  \u03a3\u2096 contribution(part k)  +  constant",
    "Each contribution is a true share of the real logit \u2014 together they reproduce the "
    "model's own output with no residual."],
   12.5, False, INK, line=1.28)
card(s, 7.00, y, 2.78, 1.3, "reconstruction error",
     f"{R['instrumentation_max_logit_error']:.0e}", "instrumented vs untouched", GREEN, 17)
card(s, 9.93, y, 2.78, 1.3, "additivity", "exact", "verified 1e-16 in float64", GREEN, 17)
tb(s, 7.00, y + 1.48, 5.71, 2.0,
   ["Branch split = Shapley values over {S, T}:",
    "- v(\u2205)=0.5, v({S})=s_head P, v({T})=t_head P, v({S,T})=final P",
    "- \u03c6_S + \u03c6_T = P_final \u2212 0.5 holds exactly",
    "- the spatial head provably never reads the temporal branch, so a branch "
    "ablation is exact for this checkpoint"],
   11.5, False, INK, line=1.24)
tb(s, MARGIN, y + 3.62, CW, 1.3,
   ["What this yields for every 3-second clip:",
    "- spatial vs temporal percentage share        - contribution of lips / skin / eyes / nose",
    "- a 16\u00d716 temporal evidence map                - per-layer contribution across 24 CLIP layers"],
   12, False, INK, line=1.26)

# ===================================================================== 4 verdicts
s = slide()
y = title(s, "The verdicts", "60 overlapping 3-second clips per video", "result 1")
b, _, _ = pic_fit(s, "f2_verdict_timeline.png", MARGIN, y, CW, 3.55)
y2 = table(s, MARGIN, b + 0.14, CW, [
    ["video", "truth", "verdict", "mean P(FAKE)", "range", "clips correct"],
    ["000.mp4", "REAL", V_R["verdict"] + "  \u2713", f"{V_R['prob']['final_mean']:.3f}",
     f"{V_R['prob']['final_min']:.2f} \u2013 {V_R['prob']['final_max']:.2f}",
     f"{V_R['clip_level_accuracy']:.0f}%   (53/60)"],
    ["000_003.mp4", "FAKE", V_F["verdict"] + "  \u2713", f"{V_F['prob']['final_mean']:.3f}",
     f"{V_F['prob']['final_min']:.2f} \u2013 {V_F['prob']['final_max']:.2f}",
     f"{V_F['clip_level_accuracy']:.0f}%   (60/60)"],
], col_w=[2.2, 1, 1.3, 1.5, 1.5, 1.8], row_h=0.31)
tb(s, MARGIN, y2 + 0.12, CW, 0.8,
   ["- The fake is detected confidently with zero disagreement across all 60 clips.",
    "- The real video stays below threshold but much closer to it (0.42), and 7 of 60 clips "
    "drift over \u2014 the model is more certain about fakes than about reals."],
   12, False, INK, line=1.26)

# ===================================================================== 5 branch share
s = slide()
y = title(s, "Spatial vs temporal: the headline answer",
          "Share of the evidence that moved each decision away from chance", "result 2")
b, _, _ = pic_fit(s, "f3_branch_share.png", MARGIN, y, CW, 3.45)
y2 = table(s, MARGIN, b + 0.12, CW, [
    ["video", "verdict", "SPATIAL", "TEMPORAL", "dominant branch", "clips where S wins"],
    ["000.mp4", "REAL", pct(V_R["branch_shapley"]["pct_spatial_mean"]),
     pct(V_R["branch_shapley"]["pct_temporal_mean"]), "TEMPORAL",
     f"{V_R['branch_shapley']['clips_spatial_dominant']}/60"],
    ["000_003.mp4", "FAKE", pct(V_F["branch_shapley"]["pct_spatial_mean"]),
     pct(V_F["branch_shapley"]["pct_temporal_mean"]), "balanced (S slightly ahead)",
     f"{V_F['branch_shapley']['clips_spatial_dominant']}/60"],
], col_w=[1.9, 1, 1.3, 1.4, 2.3, 1.8], row_h=0.31)
tb(s, MARGIN, y2 + 0.12, CW, 0.9,
   [f"- Saying FAKE is a JOINT decision: spatial {pct(V_F['branch_shapley']['pct_spatial_mean'])} "
    f"vs temporal {pct(V_F['branch_shapley']['pct_temporal_mean'])}, with the branches agreeing on "
    f"{D['counterfactual']['000_003']['clip_branch_agreement_pct']:.0f}% of clips.",
    f"- Saying REAL is mostly TEMPORAL ({pct(V_R['branch_shapley']['pct_temporal_mean'])}): the "
    f"evidence for authenticity is consistent motion, not clean-looking frames."],
   12, False, INK, line=1.26)

# ===================================================================== 6 parts
s = slide()
y = title(s, "Why it said YES (fake)",
          "The spatial branch broken into its four guided facial components", "result 3")
b, _, _ = pic_fit(s, "f4_facial_parts.png", MARGIN, y, CW, 3.25)
pf, pr = V_F["spatial_parts"], V_R["spatial_parts"]
rows = [["component", "FAKE contrib.", "share", "REAL contrib.", "share", "change (F\u2212R)"]]
for p in ("eyes", "skin", "nose", "lips"):
    rows.append([p, f"{pf['contrib_mean'][p]:+.3f}", pct(pf['share_pct'][p]),
                 f"{pr['contrib_mean'][p]:+.3f}", pct(pr['share_pct'][p]),
                 f"{D['delta_parts'][p]:+.3f}"])
y2 = table(s, MARGIN, b + 0.10, CW, rows,
           col_w=[1.7, 2.0, 1.2, 2.0, 1.2, 1.8], row_h=0.275)
tb(s, MARGIN, y2 + 0.10, CW, 0.85,
   [f"- MAIN CAUSE of the FAKE verdict: the EYES \u2014 {pct(pf['share_pct']['eyes'])} of the "
    f"spatial evidence, then skin {pct(pf['share_pct']['skin'])}.",
    f"- In the real video the nose votes REAL ({pr['contrib_mean']['nose']:+.3f}); after the swap "
    f"every component turns towards FAKE, the nose shifting most ({D['delta_parts']['nose']:+.3f})."],
   12, False, INK, line=1.26)

# ===================================================================== 7 temporal maps
s = slide()
y = title(s, "Where in the face the temporal evidence lies",
          "The temporal branch output is a 16\u00d716 grid, so it maps straight back onto the face",
          "result 4")
b, _, _ = pic_fit(s, "f5_temporal_maps.png", MARGIN, y, CW, 3.3)
yy = b + 0.14
tb(s, MARGIN, yy, 5.85, 1.9,
   ["Reading the maps:",
    "- red = this patch pushed towards FAKE, blue = towards REAL;",
    "- green dots are the averaged 68 facial landmarks;",
    "- the right panel is FAKE minus REAL \u2014 exactly where the swap created new evidence."],
   11.5, False, INK, line=1.24)
y3 = table(s, 6.75, yy, 5.96, [
    ["zone (area-normalised)", "REAL", "FAKE", "change"],
    ["eyes", f"{D['zone_density']['000']['eyes']:.4f}",
     f"{D['zone_density']['000_003']['eyes']:.4f}", f"{D['zone_density_delta']['eyes']:+.4f}"],
    ["nose", f"{D['zone_density']['000']['nose']:.4f}",
     f"{D['zone_density']['000_003']['nose']:.4f}", f"{D['zone_density_delta']['nose']:+.4f}"],
    ["mouth", f"{D['zone_density']['000']['mouth']:.4f}",
     f"{D['zone_density']['000_003']['mouth']:.4f}", f"{D['zone_density_delta']['mouth']:+.4f}"],
    ["jaw / contour", f"{D['zone_density']['000']['jaw/contour']:.4f}",
     f"{D['zone_density']['000_003']['jaw/contour']:.4f}",
     f"{D['zone_density_delta']['jaw/contour']:+.4f}"],
], col_w=[2.3, 1.1, 1.1, 1.2], row_h=0.27)
tb(s, 6.75, y3 + 0.10, 5.96, 0.6,
   f"Biggest gains after the swap: {', '.join(D['zone_density_delta_top'])} \u2014 the same "
   f"inner-face region the swap replaced.", 11.5, False, INK, line=1.22)

# ===================================================================== 8 zones + attention
s = slide()
y = title(s, "Zone attribution and component attention",
          "Area-normalising removes the advantage of physically large regions", "result 5")
b, _, _ = pic_fit(s, "f6_zones.png", MARGIN, y, CW, 2.75)
yy = b + 0.12
tb(s, MARGIN, yy, 2.25, 2.6,
   ["Sanity check on FCG:",
    "- each of the four queries lands on its own region,",
    "- so the component percentages are anatomically meaningful,",
    "- not four anonymous channels."],
   11, False, INK, line=1.22)
pic_fit(s, "f8_attention.png", 3.05, yy, 9.66, BOTTOM - yy)

# ===================================================================== 9 counterfactual
s = slide()
y = title(s, "Is one branch enough?",
          "Inference-time ablation of the released checkpoint", "result 6")
b, _, _ = pic_fit(s, "f9_counterfactual.png", MARGIN, y, 8.0, 3.1)
cf_r, cf_f = D["counterfactual"]["000"], D["counterfactual"]["000_003"]
card(s, 8.85, y, 3.86, 1.28, "fake sample", "both sufficient",
     f"S alone {cf_f['spatial_only_mean']:.3f}  \u00b7  T alone {cf_f['temporal_only_mean']:.3f}",
     GREEN, 16)
card(s, 8.85, y + 1.46, 3.86, 1.28, "real sample", "both sufficient",
     f"S alone {cf_r['spatial_only_mean']:.3f}  \u00b7  T alone {cf_r['temporal_only_mean']:.3f}",
     BLUE, 16)
tb(s, 8.85, y + 2.92, 3.86, 0.9,
   ["Per-clip branch agreement:",
    f"- fake: {cf_f['clip_branch_agreement_pct']:.0f}%     - real: {cf_r['clip_branch_agreement_pct']:.0f}%"],
   11.5, False, INK, line=1.22)
tb(s, MARGIN, b + 0.18, 8.0, 1.6,
   ["- On clean video either branch alone already reaches the correct verdict \u2014 the two are "
    "redundant here, not complementary.",
    "- Fusion still adds confidence: the fake rises 0.91 \u2192 0.93.",
    "- That redundancy is exactly what breaks under degradation (next slide)."],
   11.5, False, INK, line=1.26)

# ===================================================================== 10 layers
s = slide()
y = title(s, "Where in the frozen encoder the evidence comes from",
          "Per-layer contributions across the 24 CLIP ViT-L/14 layers", "result 7")
b, _, _ = pic_fit(s, "f7_layers.png", MARGIN, y, CW, 3.6)
tb(s, MARGIN, b + 0.18, CW, 1.7,
   [f"- Spatial evidence concentrates in DEEP layers (top for the fake: "
    f"{', '.join(str(i) for i in V_F['layers']['spatial_top'][:3])}) \u2014 semantic, "
    f"identity-level features.",
    f"- Temporal evidence also draws on SHALLOW layers (top: "
    f"{', '.join(str(i) for i in V_F['layers']['temporal_top'][:3])}) \u2014 low-level texture "
    f"and motion cues.",
    "- This is why a frozen backbone suffices: the two branches read genuinely different depths "
    "of the same encoder, which is the paper's core design claim."],
   12, False, INK, line=1.3)

# ===================================================================== 11 robustness
s = slide()
y = title(s, "Stress test on the authors' own protocol",
          "7 perturbation types \u00d7 3 severity levels, same sample pair", "result 8")
b, _, _ = pic_fit(s, "f10_robustness.png", MARGIN, y, CW, 3.35)
yy = b + 0.14
card(s, MARGIN, yy, 2.92, 1.22, "false positives", f"{RS['false_positives']} / 21",
     "REAL called FAKE", RED, 17)
card(s, 3.68, yy, 2.92, 1.22, "false negatives", f"{RS['false_negatives']} / 21",
     "FAKE called REAL", ORNG, 17)
card(s, 6.74, yy, 2.92, 1.22, "blur / compression", "always wrong",
     "GB, JPEG, VC at every level", RED, 17)
card(s, 9.79, yy, 2.92, 1.22, "block noise", "fully robust",
     "BW correct at all levels", GREEN, 17)
tb(s, MARGIN, yy + 1.36, CW, 0.9,
   ["- The failure is strongly ASYMMETRIC: degradation pushes scores UP, so the model mistakes "
    "low quality for manipulation.",
    "- Blur, JPEG-style resampling and H.264 compression flip the real video at every severity; "
    "Gaussian noise does the opposite and hides the fake."],
   11.5, False, INK, line=1.26)

# ===================================================================== 12 robustness table
s = slide()
y = title(s, "Robustness detail", "Per-condition verdicts", "result 8 \u00b7 table")
rows = [["condition", "param", "REAL P", "verdict", "FAKE P", "verdict", "error"]]
rows.append(["clean", "\u2014", f"{RS['clean_real_p']:.3f}", "REAL",
             f"{RS['clean_fake_p']:.3f}", "FAKE", "\u2014"])
for r in RS["rows"]:
    if not r["real_ok"] and not r["fake_ok"]:
        err = "both"
    elif not r["real_ok"]:
        err = "false positive"
    elif not r["fake_ok"]:
        err = "false negative"
    else:
        err = "\u2014"
    rows.append([f"{r['type']}  L{r['level']}", str(r["param"]),
                 f"{r['real_p']:.3f}", "FAKE" if r["real_p"] >= .5 else "REAL",
                 f"{r['fake_p']:.3f}", "FAKE" if r["fake_p"] >= .5 else "REAL", err])
y2 = table(s, MARGIN, y, CW, rows, col_w=[1.5, .9, 1.1, 1.1, 1.1, 1.1, 1.7],
           fs=9, head_fs=9.5, row_h=0.205)
tb(s, MARGIN, y2 + 0.12, CW, 0.5,
   "CS colour saturation \u00b7 CC colour contrast \u00b7 BW block-wise noise \u00b7 "
   "GNC Gaussian noise \u00b7 GB Gaussian blur \u00b7 JPEG down/up-sampling \u00b7 VC H.264 compression",
   10, False, MUTE)

# ===================================================================== 13 conclusion
s = slide()
y = title(s, "Answers to the question we set out to ask", None, "conclusion")
card(s, MARGIN, y, 5.95, 2.15,
     f"When it says FAKE   (P = {V_F['prob']['final_mean']:.3f})",
     f"spatial {pct(V_F['branch_shapley']['pct_spatial_mean'])} / temporal "
     f"{pct(V_F['branch_shapley']['pct_temporal_mean'])}",
     f"MAIN CAUSE: the EYES \u2014 {pct(V_F['spatial_parts']['share_pct']['eyes'])} of the spatial "
     f"evidence, then skin {pct(V_F['spatial_parts']['share_pct']['skin'])}.\n"
     f"Temporal evidence appears in the inner face "
     f"({', '.join(D['zone_density_delta_top'][:2])}) \u2014 the region the swap replaced.\n"
     f"All 60 clips agree; either branch alone suffices.", ORNG, 15)
card(s, 6.76, y, 5.95, 2.15,
     f"When it says REAL   (P = {V_R['prob']['final_mean']:.3f})",
     f"temporal {pct(V_R['branch_shapley']['pct_temporal_mean'])} / spatial "
     f"{pct(V_R['branch_shapley']['pct_spatial_mean'])}",
     f"MAIN CAUSE: temporal consistency \u2014 motion is coherent across frames.\n"
     f"Within the spatial branch the nose is the strongest REAL vote "
     f"({V_R['spatial_parts']['contrib_mean']['nose']:+.3f}, "
     f"{pct(V_R['spatial_parts']['share_pct']['nose'])}).\n"
     f"Lower margin: 7 of 60 clips cross the threshold.", BLUE, 15)
tb(s, MARGIN, y + 2.36, CW, 2.5,
   ["What this tells us about the model",
    "- Fake detection is driven by the EYES; authenticity is confirmed by MOTION CONSISTENCY \u2014 "
    "the two verdicts are not mirror images.",
    "- The spatial branch reads deep semantic layers, the temporal branch also reads shallow "
    "texture layers, so the frozen encoder is genuinely used at two depths.",
    "- On clean video the branches are redundant, so fusion adds confidence rather than coverage.",
    f"- The real weakness is QUALITY, not forgery type: {RS['false_positives']}/21 degraded "
    f"conditions raised a false alarm, because blur and compression push scores towards FAKE.",
    "- Practical consequence: a quality-aware threshold or compression augmentation would help "
    "more than an extra branch."],
   12, False, INK, line=1.3)

# ===================================================================== 14 validity
s = slide()
y = title(s, "Method validity and limitations", None, "appendix")
tb(s, MARGIN, y, 5.95, 4.4,
   ["Why these numbers can be trusted",
    f"- The instrumented model reproduces the untouched model's logits to "
    f"{R['instrumentation_max_logit_error']:.0e} \u2014 we are explaining the real checkpoint.",
    "- Contributions form an exact additive split, verified to 1e-16 in float64; they sum to the "
    "model's own logit with no residual.",
    "- Shapley efficiency holds exactly: \u03c6_S + \u03c6_T = P_final \u2212 0.5.",
    "- All 569 checkpoint tensors load with 0 missing and 0 unexpected keys; every trainable "
    "weight comes from the released file.",
    "- Labels verified from FaceForensics++ naming and confirmed visually: the difference between "
    "the pair is confined to the inner face, 22.7\u00d7 the background difference."],
   11.5, False, INK, line=1.27)
tb(s, 6.76, y, 5.95, 4.4,
   ["Limitations stated honestly",
    "- Two videos only. These are the samples shipped with the repository; FaceForensics++, "
    "Celeb-DF and DFDC are EULA-gated and cannot be downloaded without a signed agreement, so no "
    "cross-dataset generalisation is claimed.",
    "- Single-branch numbers are inference-time ablations of this checkpoint, not models "
    "re-trained with op_mode=[S] or [T], which would converge to different weights.",
    "- Patch-to-zone mapping uses nearest-landmark assignment on a 16\u00d716 grid, so zone "
    "boundaries are approximate even though patch contributions are exact.",
    "- The robustness sweep uses 8 clips per condition; the main analysis uses 60."],
   11.5, False, INK, line=1.27)

# ===================================================================== 15 reproduce
s = slide()
y = title(s, "Reproducing this analysis", None, "appendix")
tb(s, MARGIN, y, CW, 4.5,
   ["Pipeline",
    "1.  landmarks + bboxes            python -m src.preprocess.fetch_landmark_bbox ...",
    "2.  aligned face cropping         python -m src.preprocess.crop_main_face ...",
    "3.  exactness self-test           python xai/smoke.py",
    "4.  main analysis, 60 clips/video python xai/run_analysis.py --max-clips 60",
    "5.  figures + derived metrics     python xai/make_figures.py",
    "6.  robustness sweep              python xai/run_robustness.py --levels 1 3 5",
    "7.  audits                        python xai/audit/audit_math2.py , audit_pipeline.py",
    "8.  deck                          python xai/build_deck.py",
    "",
    "Environment: Windows, RTX 3050 Ti 4 GB (peak 2.5 GB at batch size 2), torch 2.6 + CUDA 12.4.",
    "Outputs: xai/out/{report,derived,robustness}.json + arrays.npz; figures in xai/figs/.",
    "Checkpoint: released FFG ViT-L/14 weights from the authors' Google Drive link."],
   12, False, INK, line=1.3)
tb(s, MARGIN, 6.62, CW, 0.45,
   "Han, Huang, Hua, Chen \u2014 Towards More General Video-based Deepfake Detection through Facial "
   "Component Guided Adaptation for Foundation Model, CVPR 2025  \u00b7  arXiv:2404.05583",
   10, False, MUTE)

path = "DFD_FCG_Explainability.pptx"
prs.save(path)
print("saved", path, "| slides:", len(prs.slides._sldIdLst))
